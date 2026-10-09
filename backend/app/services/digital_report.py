"""Read-only exports of an issued snapshot; never query live inspections."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from pathlib import PurePosixPath
from types import SimpleNamespace
import re
import json
from functools import lru_cache
from pathlib import Path
from lxml import etree

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.report_snapshot import POSITION_FIELDS, VISIT_FIELDS
from app.services.position_labels import position_label

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def _word_text(node):
    return re.sub(r'\s+', ' ', ''.join(node.xpath('.//w:t/text()', namespaces=NS))).strip()


def archived_layout(path, snapshot):
    if not path or not path.is_file():
        return []
    stat = path.stat()
    return [dict(item) for item in _archived_layout(str(path), stat.st_mtime_ns, stat.st_size, json.dumps(snapshot, sort_keys=True))]


@lru_cache(maxsize=16)
def _archived_layout(path, _modified, _size, snapshot_json):
    snapshot = json.loads(snapshot_json)
    candidates = {}
    for visit in snapshot.get('visits', []):
        for pos in visit.get('positions', []):
            remarks = position_label(SimpleNamespace(**pos)) + (' — ' + pos['inspector_notes'] if pos.get('inspector_notes') else '')
            identity = (str(visit.get('tower')), re.sub(r'\s+', ' ', remarks).strip())
            candidates.setdefault(identity, []).append(f"{visit['id']}:{pos['id']}")
    try:
        with ZipFile(path) as archive:
            root = etree.fromstring(archive.read('word/document.xml'), etree.XMLParser(resolve_entities=False, no_network=True))
        result = []
        pending = {}
        for table_index, table in enumerate(root.xpath('./w:body/w:tbl', namespaces=NS)):
            rows = table.xpath('./w:tr', namespaces=NS)
            text = _word_text(table)
            if rows and 'Insulator Details' in text and 'Tower No.' in text:
                number = re.search(r'No\.\s*(\d+)', _word_text(rows[0]))
                if not number:
                    continue
                tower = ''
                for row in rows:
                    cells = row.xpath('./w:tc', namespaces=NS)
                    if len(cells) > 1 and _word_text(cells[0]) == 'Tower No.':
                        tower = _word_text(cells[1])
                item = dict(table_index=table_index, number=int(number.group(1)), tower=tower, key=None, text=text)
                result.append(item)
                pending[number.group(1)] = item
            elif ('Max. Temp.' in text or 'Max Temp' in text) and 'Remarks' in text:
                for row_index, row in enumerate(rows):
                    cells = row.xpath('./w:tc', namespaces=NS)
                    if len(cells) < 3:
                        continue
                    item = pending.get(_word_text(cells[0]))
                    matches = candidates.get((_word_text(cells[1]), _word_text(cells[-1])), [])
                    if item and len(matches) == 1 and item['tower'] == _word_text(cells[1]):
                        item['key'] = matches[0]
                        if len(cells) == 8:
                            item.update(measurement_table_index=table_index, measurement_row_index=row_index, measurement={name: _word_text(cell) for name, cell in zip(['finding_number', 'tower', 'tmax_c', 'tref_c', 'delta_t', 'severity', 'load_current', 'remarks'], cells)})
                pending = {}
        # A snapshot key must identify exactly one original finding table.
        counts = {}
        for item in result:
            counts[item['key']] = counts.get(item['key'], 0) + 1
        for item in result:
            if counts[item['key']] != 1:
                item['key'] = None
        return result
    except (OSError, ValueError, KeyError, BadZipFile, etree.XMLSyntaxError):
        return []


def archived_finding_document(path, snapshot, keys, notice=None, overview=False, measurements=False):
    """Copy original OOXML tables/styles/image bytes; never reconstruct their appearance."""
    if not path or not path.is_file():
        raise ValueError('The archived Word document is unavailable')
    if not overview:
        selected_rows(snapshot, keys)
    original_layout = archived_layout(path, snapshot)
    layout = {item['key']: item for item in original_layout if item['key']}
    ordered = list(dict.fromkeys(keys))
    if not overview and (not ordered or any(key not in layout for key in ordered)):
        raise ValueError('Some findings cannot be matched safely to the archived Word report. View the issued Word document.')
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    stream = BytesIO()
    with ZipFile(path) as original:
        root = etree.fromstring(original.read('word/document.xml'), parser)
        body = root.find('w:body', NS)
        tables = body.findall('w:tbl', NS)
        if overview:
            finding_tables = {tables[item['table_index']] for item in original_layout}
            selected = [child for child in body if child not in finding_tables and child.tag != '{'+NS['w']+'}sectPr' and (child.tag != '{'+NS['w']+'}p' or _word_text(child) or child.xpath('.//w:drawing', namespaces=NS))]
        elif measurements:
            from copy import deepcopy
            if any('measurement' not in layout[key] for key in ordered):
                raise ValueError('Some measurements cannot be matched safely to the issued report')
            selected = []
            # Copy original headers, cell colours and row contents in requested order.
            for key in ordered:
                item = layout[key]
                source = tables[item['measurement_table_index']]
                if not selected or selected[-1][0] != item['measurement_table_index']:
                    table = deepcopy(source)
                    for row in table.findall('w:tr', NS)[1:]:
                        table.remove(row)
                    selected.append((item['measurement_table_index'], table))
                selected[-1][1].append(deepcopy(source.findall('w:tr', NS)[item['measurement_row_index']]))
            selected = [table for _, table in selected]
        else:
            selected = [tables[layout[key]['table_index']] for key in ordered]
        section = body.find('w:sectPr', NS)
        for child in list(body):
            body.remove(child)
        if notice:
            paragraph = etree.SubElement(body, '{'+NS['w']+'}p')
            run = etree.SubElement(paragraph, '{'+NS['w']+'}r')
            etree.SubElement(run, '{'+NS['w']+'}t').text = notice
        for node in selected:
            body.append(node)
            # Separate findings without altering the contents of the original tables.
            etree.SubElement(body, '{'+NS['w']+'}p')
        if section is not None:
            body.append(section)
        used = set(root.xpath('//@r:embed | //@r:id | //@r:link', namespaces=NS))
        relations = etree.fromstring(original.read('word/_rels/document.xml.rels'), parser)
        keep_media = set()
        names = set(original.namelist())
        for rel in list(relations):
            if rel.get('Type', '').endswith('/image'):
                if rel.get('Id') not in used:
                    relations.remove(rel)
                else:
                    keep_media.add(str(PurePosixPath('word') / rel.get('Target')))
        # Keep images used by headers, footers and other retained document parts.
        for name in names:
            if name.endswith('.rels') and name != 'word/_rels/document.xml.rels':
                rels = etree.fromstring(original.read(name), parser)
                parent = PurePosixPath(name).parent.parent
                for rel in rels:
                    if rel.get('Type', '').endswith('/image') and rel.get('TargetMode') != 'External':
                        keep_media.add(str(parent / rel.get('Target')))
        excluded = {name for name in names if name.startswith('word/media/') and name not in keep_media}
        content_types = etree.fromstring(original.read('[Content_Types].xml'), parser)
        for entry in list(content_types):
            if entry.get('PartName', '').lstrip('/') in excluded:
                content_types.remove(entry)
        replacements = {'word/document.xml': root, 'word/_rels/document.xml.rels': relations, '[Content_Types].xml': content_types}
        with ZipFile(stream, 'w') as output:
            for info in original.infolist():
                if info.filename not in excluded:
                    data = etree.tostring(replacements[info.filename], xml_declaration=True, encoding='UTF-8', standalone=True) if info.filename in replacements else original.read(info.filename)
                    output.writestr(info, data)
    return stream.getvalue()


def archived_evidence_index(path, snapshot):
    if not path or not path.is_file():
        return {}
    stat = path.stat()
    return dict(_archived_evidence_index(str(path), stat.st_mtime_ns, stat.st_size, json.dumps(snapshot, sort_keys=True)))


@lru_cache(maxsize=16)
def _archived_evidence_index(path_string, _modified, _size, snapshot_json):
    """Associate original template photos only when the saved finding matches uniquely.

    Measurement rows connect printed finding numbers to frozen tower/position labels.
    A repeated/ambiguous label or a skipped corrupt photo fails closed, rather than
    attributing another finding's photograph. No live Image rows are used.
    """
    path = Path(path_string)
    snapshot = json.loads(snapshot_json)
    def text(node):
        return re.sub(r'\s+', ' ', ''.join(node.xpath('.//w:t/text()', namespaces=NS))).strip()
    candidates = {}
    for visit in snapshot.get('visits', []):
        for pos in visit.get('positions', []):
            remarks = position_label(SimpleNamespace(**pos)) + (' — ' + pos['inspector_notes'] if pos.get('inspector_notes') else '')
            identity = (str(visit.get('tower')), re.sub(r'\s+', ' ', remarks).strip())
            candidates.setdefault(identity, []).append((f"{visit['id']}:{pos['id']}", pos))
    try:
        with ZipFile(path) as archive:
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            root = etree.fromstring(archive.read('word/document.xml'), parser)
            rels = etree.fromstring(archive.read('word/_rels/document.xml.rels'), parser)
            targets = {}
            for rel in rels:
                target = rel.get('Target', '')
                member = str(PurePosixPath('word') / target)
                if rel.get('TargetMode') != 'External' and member.startswith('word/media/') and '..' not in PurePosixPath(member).parts and member in archive.namelist():
                    targets[rel.get('Id')] = member
        pending = {}
        result = {}
        for table in root.xpath('./w:body/w:tbl', namespaces=NS):
            table_rows = table.xpath('./w:tr', namespaces=NS)
            if not table_rows:
                continue
            content = text(table)
            if 'Insulator Details' in content and 'Tower No.' in content:
                match = re.search(r'No\.\s*(\d+)', text(table_rows[0]))
                if match:
                    pending[match.group(1)] = table
            elif ('Max. Temp.' in content or 'Max Temp' in content) and 'Remarks' in content:
                for row in table_rows:
                    cells = row.xpath('./w:tc', namespaces=NS)
                    if len(cells) < 3:
                        continue
                    seq = text(cells[0])
                    matches = candidates.get((text(cells[1]), text(cells[-1])), [])
                    finding = pending.get(seq)
                    if len(matches) != 1 or finding is None:
                        continue
                    key, position = matches[0]
                    for caption, category in [('Thermal (Full):', 'TH Full'), ('Thermal (Close):', 'TH Close'), ('Visual (Full):', 'RGB Full'), ('Visual (Close):', 'RGB Close')]:
                        image_cells = [cell for cell in finding.xpath('.//w:tc', namespaces=NS) if caption in text(cell)]
                        ids = [item for cell in image_cells for item in cell.xpath('.//a:blip/@r:embed', namespaces=NS)]
                        evidence = [(index, item) for index, item in enumerate(position.get('evidence', [])) if item['image_type'] == category]
                        if len(ids) == len(evidence) and all(rid in targets for rid in ids):
                            for (index, _), rid in zip(evidence, ids):
                                result[f'{key}:{index}'] = targets[rid]
                pending = {}
        return result
    except (OSError, ValueError, KeyError, BadZipFile, etree.XMLSyntaxError):
        return {}


def snapshot_rows(snapshot):
    return [{
        'key': f"{visit['id']}:{position['id']}",
        'tower': visit.get('tower'), 'area': visit.get('area'),
        'line_sector': visit.get('line_sector'),
        **{field: visit.get(field) for field in VISIT_FIELDS},
        **{field: position.get(field) for field in POSITION_FIELDS},
        'evidence_count': len(position.get('evidence', [])),
    } for visit in snapshot.get('visits', []) for position in visit.get('positions', [])]


def selected_rows(snapshot, keys):
    rows = snapshot_rows(snapshot)
    selected = set(keys)
    if selected - {row['key'] for row in rows}:
        raise ValueError('Selection contains findings outside this issued report')
    return [row for row in rows if row['key'] in selected]


def export_snapshot(snapshot, number, issued, keys, kind, archive_path=None):
    rows = selected_rows(snapshot, keys)
    if not rows:
        raise ValueError('Select at least one finding')
    note = 'Filtered inspection extract from the issued report. This is not a newly approved report. Original photographs and template remain in the archived Word document.'
    stream = BytesIO()
    fields = [key for key in rows[0] if key != 'key']
    if kind == 'xlsx':
        book = Workbook()
        sheet = book.active
        sheet.title = 'Findings'
        sheet.append([field.replace('_', ' ').title() for field in fields])
        for row in rows:
            sheet.append([row[field] for field in fields])
        # Explicit text type prevents spreadsheet formulas in inspection notes.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='10485B')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for index, field in enumerate(fields, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 36 if 'notes' in field else 22
        provenance = book.create_sheet('Report source')
        for label, value in [('Report number', number), ('Issued at', str(issued)), ('Team', snapshot.get('team_name')), ('Selected findings', len(rows)), ('Total findings', len(snapshot_rows(snapshot))), ('Notice', note)]:
            provenance.append([label, value])
        for row in provenance:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        provenance.column_dimensions['A'].width = 24
        provenance.column_dimensions['B'].width = 100
        book.save(stream)
    else:
        doc = Document()
        section = doc.sections[0]
        section.left_margin = section.right_margin = Inches(.7)
        doc.add_heading('Filtered inspection extract', 0)
        doc.add_paragraph(f'{number} | Issued: {issued} | {len(rows)} of {len(snapshot_rows(snapshot))} findings')
        doc.add_paragraph(note)
        evidence = archived_evidence_index(archive_path, snapshot)
        for index, row in enumerate(rows, 1):
            doc.add_heading(f"{index}. Tower {row['tower']} — {row.get('position_code') or 'Finding'}", 1)
            table = doc.add_table(rows=0, cols=2)
            table.style = 'Light Shading Accent 1'
            for field in fields:
                if row[field] is not None and row[field] != '':
                    cells = table.add_row().cells
                    cells[0].text = field.replace('_', ' ').title()
                    cells[1].text = str(row[field])
            if archive_path and evidence:
                with ZipFile(archive_path) as archive:
                    for evidence_key, member in evidence.items():
                        if evidence_key.rsplit(':', 1)[0] == row['key']:
                            doc.add_picture(BytesIO(archive.read(member)), width=Inches(3.5))
        doc.save(stream)
    return stream.getvalue()
