"""Groups the official OETC-format report (services/oetc_report.py) across more than one team's
campaign into a single .docx — by Area (every team currently working that area, in the date range)
or, for the fully "collected" version, by every area in the catalog, in turn.

Each team's own section is rendered completely unchanged from the single-team report (same
customer template, same layout, same page design) — grouping only decides which team-sections go
into the file and in what order: Area, then Team, then — already handled inside each team's own
section by services/oetc_report.py — Mission. Nothing about the customer's official page is ever
touched; sections are just concatenated with a page break between them via docxcompose, the
standard library for composing several Word documents into one.
"""
from __future__ import annotations

import datetime as dt
import io
import re

from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer
from sqlalchemy.orm import Session, joinedload

from app.models import Area, Position, Team, Tower, Visit
from app.schemas import LineInspectionReportRequest, OetcAreaReportRequest, OetcConsolidatedReportRequest
from app.services.oetc_report import render_oetc_line_report_docx


class GroupedReportBlock:
    """One team's section within the merged file — everything the router needs to persist a
    LineInspectionReport row (and its own saved .docx + linked images) for it afterwards, without
    re-deriving or re-rendering anything."""

    __slots__ = ("area", "team", "visits", "report_number", "docx_bytes")

    def __init__(self, area: str, team: Team, visits: list[Visit], report_number: str, docx_bytes: bytes):
        self.area = area
        self.team = team
        self.visits = visits
        self.report_number = report_number
        # This exact section's own standalone .docx (pre-merge) — the same bytes composed into the
        # combined file, saved separately so this one team's report can be traced/redownloaded on
        # its own later, same as a plain single-team report.
        self.docx_bytes = docx_bytes


def _slug(text: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "-", text or "").strip("-").upper()
    return s or "X"


def _visits_for_area(db: Session, area: str, start_date: dt.date, end_date: dt.date) -> list[Visit]:
    return (
        db.query(Visit)
        .join(Tower, Visit.tower_id == Tower.id)
        .options(joinedload(Visit.positions).joinedload(Position.images), joinedload(Visit.tower), joinedload(Visit.team))
        .filter(
            Tower.area == area,
            Visit.team_id.isnot(None),
            Visit.inspection_date.isnot(None),
            Visit.inspection_date >= start_date,
            Visit.inspection_date <= end_date,
        )
        .all()
    )


def _group_by_team(visits: list[Visit]) -> list[tuple[Team, list[Visit]]]:
    """Every distinct team present, each with its own visits, sorted by team name — the "by team"
    half of "by area, by team, by mission" (the "by mission" half already happens inside
    build_oetc_line_report_context, which sorts a team's own visits by mission_seq/date)."""
    by_team: dict[int, list[Visit]] = {}
    for v in visits:
        by_team.setdefault(v.team_id, []).append(v)
    groups = [(visits_[0].team, visits_) for visits_ in by_team.values()]
    groups.sort(key=lambda g: g[0].name)
    return groups


def _add_page_break(doc: DocxDocument) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def _team_payload(base, team_id: int, report_number: str) -> LineInspectionReportRequest:
    return LineInspectionReportRequest(
        team_id=team_id,
        start_date=base.start_date,
        end_date=base.end_date,
        report_number=report_number,
        overall_condition=base.overall_condition,
        probable_cause=base.probable_cause,
        corrective_action=base.corrective_action,
        additional_comments=base.additional_comments,
        prepared_by=base.prepared_by,
        reviewed_by=base.reviewed_by,
        approved_by=base.approved_by,
        approval_date=base.approval_date,
    )


def _render_merged(plan: list[tuple[str, Team, list[Visit]]], base, base_number: str) -> tuple[bytes, list[GroupedReportBlock]]:
    """plan: (area, team, visits) tuples in the exact final order the merged file should have."""
    master: DocxDocument | None = None
    composer: Composer | None = None
    blocks: list[GroupedReportBlock] = []
    for area, team, visits in plan:
        sub_number = f"{base_number}-{_slug(area)}-{_slug(team.name)}"
        payload = _team_payload(base, team.id, sub_number)
        docx_bytes = render_oetc_line_report_docx(team, visits, payload)
        sub_doc = DocxDocument(io.BytesIO(docx_bytes))
        if master is None:
            master = sub_doc
            composer = Composer(master)
        else:
            _add_page_break(master)
            composer.append(sub_doc)
        blocks.append(GroupedReportBlock(area, team, visits, sub_number, docx_bytes))
    buf = io.BytesIO()
    composer.save(buf)
    return buf.getvalue(), blocks


def plan_area_report(db: Session, payload: OetcAreaReportRequest) -> list[tuple[str, Team, list[Visit]]]:
    visits = _visits_for_area(db, payload.area, payload.start_date, payload.end_date)
    return [(payload.area, team, v) for team, v in _group_by_team(visits)]


def plan_consolidated_report(db: Session, payload: OetcConsolidatedReportRequest) -> list[tuple[str, Team, list[Visit]]]:
    areas = [a.name for a in db.query(Area).order_by(Area.name).all()]
    plan: list[tuple[str, Team, list[Visit]]] = []
    for area in areas:
        visits = _visits_for_area(db, area, payload.start_date, payload.end_date)
        for team, v in _group_by_team(visits):
            plan.append((area, team, v))
    return plan


def plan_report_numbers(plan: list[tuple[str, Team, list[Visit]]], base_number: str) -> list[str]:
    """The report_number each team-section in `plan` will get, without rendering anything — lets the
    router check every one is free before spending the time to build the actual documents."""
    return [f"{base_number}-{_slug(area)}-{_slug(team.name)}" for area, team, _ in plan]


def render_area_report(db: Session, payload: OetcAreaReportRequest) -> tuple[bytes | None, list[GroupedReportBlock]]:
    plan = plan_area_report(db, payload)
    if not plan:
        return None, []
    return _render_merged(plan, payload, payload.report_number)


def render_consolidated_report(db: Session, payload: OetcConsolidatedReportRequest) -> tuple[bytes | None, list[GroupedReportBlock]]:
    plan = plan_consolidated_report(db, payload)
    if not plan:
        return None, []
    return _render_merged(plan, payload, payload.report_number)
