"""The "team activity" master report — every team's field work, sorted the way an admin actually
thinks about progress: team, then day, then which towers that day, then each string's measurements
and evidence. Built from live Visit/Position/Image data, not a separate record — a "mission" IS a
Visit with team_id set (see routers/visits.py), so this just re-shapes that same data into the
hierarchy an admin wants to skim without opening each visit individually.

Only the formal per-position checklist images (TH Full/TH Close/RGB Full/RGB Close) that actually
have content are included — the free-form visit-photos gallery (VisitPhoto) is a separate, looser
scratch space and deliberately left out here; this report is about the official evidence trail.
A position is only included once there's something to show for it: its Direction has been set
(work started) or it already carries a real image — an untouched placeholder position is noise.
"""
from __future__ import annotations

import datetime as dt
from io import BytesIO

from fastapi import HTTPException
from sqlalchemy.orm import Session, joinedload

from app.models import IMAGE_TYPE_CHOICES, OHL_CHOICES, PHASE_CHOICES, STRING_CHOICES, Image, Position, User, UserRole, Visit

_OHL_ORDER = {v: i for i, v in enumerate(OHL_CHOICES)}
_PHASE_ORDER = {v: i for i, v in enumerate(PHASE_CHOICES)}
_STRING_ORDER = {v: i for i, v in enumerate(STRING_CHOICES)}
_TYPE_ORDER = {v: i for i, v in enumerate(IMAGE_TYPE_CHOICES)}


def _position_sort_key(pos: Position):
    return (_OHL_ORDER.get(pos.ohl, 99), _PHASE_ORDER.get(pos.phase, 99), _STRING_ORDER.get(pos.string, 99))


def _image_sort_key(img: Image):
    return (_TYPE_ORDER.get(img.image_type, 99), img.sequence)


def _position_has_activity(pos: Position) -> bool:
    return bool(pos.direction) or any(img.file_path for img in pos.images)


def query_team_activity_visits(
    db: Session,
    user: User,
    team_id: int | None = None,
    start_date: dt.date | None = None,
    end_date: dt.date | None = None,
) -> list[Visit]:
    q = (
        db.query(Visit)
        .filter(Visit.team_id.isnot(None))
        .options(
            joinedload(Visit.team),
            joinedload(Visit.tower),
            joinedload(Visit.positions).joinedload(Position.images),
        )
    )
    if user.role == UserRole.TEAM_MEMBER.value:
        # A manager-level report (whole team, day by day) — not part of a team_member's narrower
        # "my assigned missions" workspace.
        raise HTTPException(status_code=403, detail="Not available for team-member accounts")
    if user.role == UserRole.TEAM_LEADER.value:
        q = q.filter(Visit.team_id == user.team_id) if user.team_id else q.filter(False)
    elif team_id:
        q = q.filter(Visit.team_id == team_id)
    if start_date:
        q = q.filter(Visit.inspection_date >= start_date)
    if end_date:
        q = q.filter(Visit.inspection_date <= end_date)
    return q.order_by(
        Visit.team_id, Visit.inspection_date.asc().nullslast(), Visit.tower_id, Visit.mission_seq
    ).all()


def build_team_activity_tree(visits: list[Visit]) -> list[dict]:
    """Groups the given visits (already filtered/ordered by query_team_activity_visits) into
    team -> day -> tower -> position -> image, in presentation order."""
    teams: dict[int, dict] = {}
    for visit in visits:
        if not visit.team:
            continue
        team_bucket = teams.setdefault(
            visit.team_id, {"team_id": visit.team_id, "team_name": visit.team.name, "days": {}}
        )
        day_key = visit.inspection_date.isoformat() if visit.inspection_date else "(no date set)"
        day_bucket = team_bucket["days"].setdefault(day_key, {"date": day_key, "towers": []})

        positions_out = []
        for pos in sorted(visit.positions, key=_position_sort_key):
            if not _position_has_activity(pos):
                continue
            images_out = [
                {
                    "id": img.id,
                    "image_type": img.image_type,
                    "image_code": img.image_code,
                    "sequence": img.sequence,
                    "capture_date": img.capture_date.isoformat() if img.capture_date else None,
                    "capture_time": img.capture_time.isoformat() if img.capture_time else None,
                    "evidence_status": img.evidence_status,
                    "annotated": bool(img.annotated_path),
                    "uploaded_at": img.uploaded_at.isoformat() if img.uploaded_at else None,
                }
                for img in sorted(pos.images, key=_image_sort_key)
                if img.file_path
            ]
            positions_out.append(
                {
                    "id": pos.id,
                    "ohl": pos.ohl,
                    "phase": pos.phase,
                    "string": pos.string,
                    "direction": pos.direction,
                    "tower_proximity": pos.tower_proximity,
                    "position_code": pos.position_code,
                    "screening_result": pos.screening_result,
                    "hotspot": pos.hotspot,
                    "images": images_out,
                }
            )

        day_bucket["towers"].append(
            {
                "visit_id": visit.id,
                "tower_id": visit.tower_id,
                "tower_code": visit.tower.tower_id,
                "area": visit.tower.area,
                "mission_seq": visit.mission_seq,
                "mission_status": visit.mission_status,
                "inspector_name": visit.inspector_name,
                "positions": positions_out,
            }
        )

    out = []
    for team_bucket in teams.values():
        days_sorted = sorted(team_bucket["days"].values(), key=lambda d: d["date"])
        out.append({"team_id": team_bucket["team_id"], "team_name": team_bucket["team_name"], "days": days_sorted})
    out.sort(key=lambda t: t["team_name"].lower())
    return out


_HEADERS = [
    "Team", "Day", "Tower", "Area", "Mission #", "Mission status",
    "OHL", "Phase", "String", "Direction", "Inner/Outer", "Position code",
    "Image type", "Image code", "Capture date", "Capture time", "Evidence status", "Annotated",
]


def build_team_activity_workbook(tree: list[dict]) -> bytes:
    """Flattens the same tree into one filterable/sortable Excel sheet — every row carries its full
    Team/Day/Tower/Position context (rather than merged cells) so Excel's own sort and AutoFilter
    keep working no matter how the admin slices it."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Team activity"

    ws.append(_HEADERS)
    header_fill = PatternFill(start_color="0F3A4D", end_color="0F3A4D", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    row_idx = 1
    for team in tree:
        for day in team["days"]:
            for tower in day["towers"]:
                if not tower["positions"]:
                    continue
                for pos in tower["positions"]:
                    base = [
                        team["team_name"], day["date"], tower["tower_code"], tower["area"] or "",
                        tower["mission_seq"] or "", tower["mission_status"],
                        pos["ohl"], pos["phase"], pos["string"], pos["direction"] or "",
                        pos["tower_proximity"] or "", pos["position_code"] or "",
                    ]
                    if not pos["images"]:
                        ws.append(base + ["", "", "", "", "", ""])
                        row_idx += 1
                        continue
                    for img in pos["images"]:
                        ws.append(
                            base
                            + [
                                img["image_type"],
                                img["image_code"] or "",
                                img["capture_date"] or "",
                                (img["capture_time"] or "")[:5],
                                img["evidence_status"],
                                "Yes" if img["annotated"] else "No",
                            ]
                        )
                        row_idx += 1

    ws.freeze_panes = "A2"
    if row_idx > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}{row_idx}"

    widths = [16, 12, 14, 12, 10, 13, 7, 7, 7, 10, 24, 11, 26, 12, 12, 16, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
