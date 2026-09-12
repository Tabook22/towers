"""Bulk tower import/update from an Excel (.xlsx) file — admin uploads a spreadsheet with one row
per tower, and it's upserted by Tower ID: an existing Tower ID gets its fields updated in place, a
new one gets created. Meant for standing up (or correcting) the tower list in one go — teams and
missions are built on top of towers, so getting a whole line's worth of them in at once, rather
than one form submission at a time, is the actual point.
"""
from __future__ import annotations

import io
import re

from sqlalchemy.orm import Session

from app.models import Tower

# Column header -> Tower field. Matching is case-insensitive and whitespace-trimmed, and column
# ORDER in the sheet doesn't matter — only the header text does — so a real-world export from
# another system just needs its columns renamed/reordered to match, not rebuilt from scratch.
COLUMN_MAP: dict[str, str] = {
    "tower id": "tower_id",
    "voltage": "voltage",
    "tower type": "tower_type",
    "area": "area",
    "line sector": "line_sector",
    "location name": "location_name",
    "height (m)": "height_m",
    "height": "height_m",
    "latitude": "latitude",
    "longitude": "longitude",
    "notes": "notes",
    # Combined GPS cell used in field exports (DMS like 17°02'32.18"N 54°29'23.82"E).
    "tower gps location": "_gps",
    "gps location": "_gps",
    "gps": "_gps",
    "coordinates": "_gps",
    "coord": "_gps",
    "tower numbers": "_tower_number",
    "tower number": "_tower_number",
}
HEADER_ROW = ["Tower ID", "Voltage", "Tower Type", "Area", "Line Sector", "Location Name", "Height (m)", "Latitude", "Longitude", "Notes"]
# Extra columns on the *export* only — the importer ignores unknown headers, so this file can be
# edited and uploaded again without those two columns wiping anything.
EXPORT_EXTRA = ["Assigned Team", "Active"]
EXPORT_HEADER = HEADER_ROW + EXPORT_EXTRA
FLOAT_FIELDS = {"height_m", "latitude", "longitude"}
FLOAT_RANGES = {"height_m": (0, 1000), "latitude": (-90, 90), "longitude": (-180, 180)}


def build_tower_import_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Towers"
    ws.append(HEADER_ROW)
    header_fill = PatternFill(start_color="0F3A4D", end_color="0F3A4D", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    ws.append(["ARSD 92", "132 kV", "Suspension (Tangent) Tower", "Dhofar", "Ittin - Thumrait", "Al Ain Corridor, Pole 14", 32.5, 17.01972, 54.08972, "Example row — edit or delete"])
    ws.freeze_panes = "A2"
    widths = [16, 12, 26, 14, 20, 26, 11, 12, 12, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    notes = wb.create_sheet("Read me")
    notes.column_dimensions["A"].width = 90
    for line in [
        "How this import works",
        "",
        "- Tower ID is the only required column — it's the unique key.",
        "- If a Tower ID already exists in the system, that tower's other fields are UPDATED to match this row.",
        "- If a Tower ID doesn't exist yet, a new tower is created.",
        "- Leave a cell blank to leave that field untouched on an existing tower (blank does not erase it).",
        "- Height is in metres. Latitude/Longitude are decimal degrees (e.g. 17.01972, 54.08972).",
        "- Column order doesn't matter — only the header text in row 1 of the 'Towers' sheet does.",
        "- Delete the example row before uploading your real list, or just overwrite it.",
    ]:
        notes.append([line])
    notes["A1"].font = Font(bold=True, size=13)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _style_header(ws) -> None:
    from openpyxl.styles import Font, PatternFill

    header_fill = PatternFill(start_color="0F3A4D", end_color="0F3A4D", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill


def export_towers_workbook(towers: list[Tower]) -> bytes:
    """Full catalog dump: same columns the importer understands, plus assigned team and active flag."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Towers"
    ws.append(EXPORT_HEADER)
    _style_header(ws)
    for t in towers:
        ws.append(
            [
                t.tower_id,
                t.voltage,
                t.tower_type,
                t.area,
                t.line_sector,
                t.location_name,
                t.height_m,
                t.latitude,
                t.longitude,
                t.notes,
                t.assigned_team.name if t.assigned_team else None,
                "Yes" if t.is_active else "No",
            ]
        )
    ws.freeze_panes = "A2"
    widths = [16, 12, 26, 14, 20, 26, 11, 12, 12, 30, 18, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    notes = wb.create_sheet("Read me")
    notes.column_dimensions["A"].width = 90
    for line in [
        "Tower catalog export",
        "",
        f"{len(towers)} tower(s) in this file, including deactivated ones (Active = No).",
        "",
        "The first 10 columns match the Import from Excel template. You can edit this sheet and",
        "upload it again: existing Tower IDs are updated, new IDs are created. Nothing is deleted.",
        "Assigned Team and Active are for reference only — they are ignored on import.",
        "Leave a cell blank on re-import to leave that field unchanged on an existing tower.",
    ]:
        notes.append([line])
    notes["A1"].font = Font(bold=True, size=13)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _normalize_header(h) -> str:
    return str(h or "").strip().lower()


# 17°02'32.18"N 54°29'23.82"E  (minutes/seconds optional; curly quotes accepted)
_DMS_PAIR = re.compile(
    r"""
    (?P<lat_d>\d+(?:\.\d+)?)\s*[°º]\s*
    (?:(?P<lat_m>\d+(?:\.\d+)?)\s*['′’]\s*)?
    (?:(?P<lat_s>\d+(?:\.\d+)?)\s*["″”]\s*)?
    (?P<lat_h>[NSns])
    \s*[,;\s]+\s*
    (?P<lon_d>\d+(?:\.\d+)?)\s*[°º]\s*
    (?:(?P<lon_m>\d+(?:\.\d+)?)\s*['′’]\s*)?
    (?:(?P<lon_s>\d+(?:\.\d+)?)\s*["″”]\s*)?
    (?P<lon_h>[EWew])
    """,
    re.VERBOSE,
)
_DEC_PAIR = re.compile(r"^\s*(-?\d+(?:\.\d+)?)\s*[,;]\s*(-?\d+(?:\.\d+)?)\s*$")
_KV = re.compile(r"^(\d+)\s*k\s*v$", re.IGNORECASE)


def _dms_to_decimal(deg, minutes, seconds, hemi: str) -> float:
    value = float(deg) + float(minutes or 0) / 60.0 + float(seconds or 0) / 3600.0
    if hemi.upper() in ("S", "W"):
        value = -value
    return value


def parse_gps(value) -> tuple[float | None, float | None]:
    """Turn a combined GPS cell into decimal (lat, lng). Accepts DMS or 'lat, lng'."""
    if value is None:
        return None, None
    text = str(value).strip()
    if not text:
        return None, None
    match = _DMS_PAIR.search(text)
    if match:
        lat = _dms_to_decimal(match["lat_d"], match["lat_m"], match["lat_s"], match["lat_h"])
        lng = _dms_to_decimal(match["lon_d"], match["lon_m"], match["lon_s"], match["lon_h"])
        if -90 <= lat <= 90 and -180 <= lng <= 180:
            return lat, lng
        return None, None
    match = _DEC_PAIR.match(text)
    if match:
        lat, lng = float(match.group(1)), float(match.group(2))
        if -90 <= lat <= 90 and -180 <= lng <= 180:
            return lat, lng
    return None, None


def _parse_float(value, field: str, row_num: int, warnings: list[str]) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        f = float(value)
    except (TypeError, ValueError):
        warnings.append(f"Row {row_num}: '{value}' isn't a number for {field} — left blank")
        return None
    lo, hi = FLOAT_RANGES[field]
    if not (lo <= f <= hi):
        warnings.append(f"Row {row_num}: {field}={f} is outside {lo}-{hi} — left blank")
        return None
    return f


def import_towers_from_excel(db: Session, raw: bytes) -> dict:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # a non-.xlsx file, corrupt upload, etc.
        return {"created": 0, "updated": 0, "total_rows": 0, "warnings": [f"Could not read that file as an Excel workbook: {exc}"]}

    # The first sheet is always the data sheet — "Towers" in the template, but don't require that
    # exact name in case someone renames or rebuilds the sheet from scratch.
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "total_rows": 0, "warnings": ["That sheet is empty."]}

    header = [_normalize_header(h) for h in rows[0]]
    col_index: dict[str, int] = {}
    for i, h in enumerate(header):
        field = COLUMN_MAP.get(h)
        if field and field not in col_index:  # first matching column wins if a name repeats
            col_index[field] = i
    if "tower_id" not in col_index:
        return {
            "created": 0,
            "updated": 0,
            "total_rows": 0,
            "warnings": ["No 'Tower ID' column found in row 1 — download the template to see the expected headers."],
        }

    created = 0
    updated = 0
    warnings: list[str] = []
    seen_ids: set[str] = set()

    for row_num, row in enumerate(rows[1:], start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue  # a blank spacer row — not an error, just skip it
        get = lambda field: row[col_index[field]] if field in col_index and col_index[field] < len(row) else None  # noqa: E731

        tower_id_raw = get("tower_id")
        tower_id = str(tower_id_raw).strip() if tower_id_raw is not None else ""
        if not tower_id:
            warnings.append(f"Row {row_num}: missing Tower ID — row skipped")
            continue
        if tower_id.lower() in seen_ids:
            warnings.append(f"Row {row_num}: duplicate Tower ID '{tower_id}' in this file — later row wins")
        seen_ids.add(tower_id.lower())

        tower = db.query(Tower).filter(Tower.tower_id.ilike(tower_id)).first()
        is_new = tower is None
        if is_new:
            tower = Tower(tower_id=tower_id)
            db.add(tower)
            # The session doesn't autoflush (see database.py) — without this, a second row for the
            # same new Tower ID later in this same file wouldn't find it via the query above and
            # would insert a duplicate instead of updating it.
            db.flush()
        else:
            tower.tower_id = tower_id  # normalize casing/whitespace to what's in the sheet

        for field in ("voltage", "tower_type", "area", "line_sector", "location_name", "notes"):
            value = get(field)
            if value is not None and str(value).strip() != "":
                text = str(value).strip()
                if field == "voltage":
                    kv = _KV.match(text.replace(" ", ""))
                    if kv:
                        text = f"{kv.group(1)} kV"
                setattr(tower, field, text)

        number = get("_tower_number")
        if (not tower.location_name) and number is not None and str(number).strip() != "":
            tower.location_name = f"Tower {str(number).strip()}"

        for field in FLOAT_FIELDS:
            value = get(field)
            if value is not None and str(value).strip() != "":
                if field in ("latitude", "longitude") and parse_gps(value)[0] is not None:
                    continue
                parsed = _parse_float(value, field, row_num, warnings)
                if parsed is not None:
                    setattr(tower, field, parsed)

        if tower.latitude is None or tower.longitude is None:
            gps_lat, gps_lng = parse_gps(get("_gps"))
            if gps_lat is None:
                # Some sheets dump the combined DMS string into the Latitude column.
                gps_lat, gps_lng = parse_gps(get("latitude"))
            if gps_lat is not None and gps_lng is not None:
                tower.latitude = gps_lat
                tower.longitude = gps_lng
            elif get("_gps") not in (None, "") or (get("latitude") not in (None, "") and tower.latitude is None):
                warnings.append(f"Row {row_num}: could not parse GPS for '{tower_id}' — location left blank")

        if is_new:
            created += 1
        else:
            updated += 1

    db.commit()
    try:
        from app.database import engine
        from app.migrations import backfill_areas_from_towers

        backfill_areas_from_towers(engine)
    except Exception:
        pass
    return {"created": created, "updated": updated, "total_rows": created + updated, "warnings": warnings}
