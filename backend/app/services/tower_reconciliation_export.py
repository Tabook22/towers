"""Coloured checklist export using the same authorised reconciliation data as the UI."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

STATUS_COLOURS = {
    'red': ('FCE4E6', '9C0006'), 'amber': ('FFF2CC', '9C6500'),
    'green': ('E2F0D9', '375623'), 'grey': ('E7E6E6', '595959'),
}
SHEET_TITLES = {
    'summary': 'Summary - الملخص', 'full': 'Tower checklist - قائمة الأبراج',
    'red': 'Red - بدون أدلة', 'amber': 'Amber - غير مكتمل',
    'green': 'Green - جاهز', 'grey': 'Grey - تاريخي',
}
ARABIC_LABELS = {
    'Team': 'الفريق', 'Tower': 'البرج', 'Scope': 'نطاق التكليف',
    'Assigned towers': 'الأبراج المكلّف بها الفريق', 'TOTAL': 'الإجمالي',
    'No archived evidence (RED)': 'لا توجد أدلة مؤرشفة (أحمر)',
    'Incomplete with evidence (AMBER)': 'توجد أدلة لكن العمل غير مكتمل (أصفر)',
    'Screening + evidence ready (GREEN)': 'الفحص والأدلة جاهزة (أخضر)',
    'Historical entries (GREY)': 'سجلات أدلة تاريخية (رمادي)',
    'Latest inspection date': 'تاريخ آخر فحص', 'Visit ID': 'رقم سجل الزيارة',
    'All-time images and field photos': 'جميع الصور المؤرشفة والصور الميدانية',
    'Installed positions': 'مواضع العوازل المركّبة', 'Screened positions': 'مواضع العوازل التي تم فحصها',
    'Expected category slots': 'خانات فئات الصور المطلوبة', 'Covered category slots': 'خانات فئات الصور المكتملة',
    'Missing category slots': 'خانات فئات الصور الناقصة', 'Screening ready': 'الفحص مكتمل',
    'Four-category coverage': 'تغطية فئات الصور الأربع', 'Screening and evidence ready': 'الفحص والأدلة جاهزة',
    'Attention': 'حالة المتابعة', 'Current assigned team': 'الفريق المكلّف حاليًا',
    'Current assignment': 'ضمن التكليف الحالي',
    'Historical evidence outside current assignment': 'أدلة تاريخية خارج التكليف الحالي',
    'No archived evidence': 'لا توجد أدلة مؤرشفة', 'Screening + evidence ready': 'الفحص والأدلة جاهزة',
    'Incomplete screening or categories': 'الفحص أو فئات الصور غير مكتملة', 'Historical evidence': 'أدلة تاريخية',
}


def _label(english):
    return english + '\n' + ARABIC_LABELS[english]


def _status(row):
    if row[2].split('\n')[0] != 'Current assignment':
        return 'grey'
    if row[5] == 0:
        return 'red'
    return 'green' if row[13] else 'amber'


def build_reconciliation_workbook(teams: list[dict]) -> bytes:
    book = Workbook()
    summary = book.active
    summary.title = SHEET_TITLES['summary']
    summary.append([_label(v) for v in ['Team', 'Assigned towers', 'No archived evidence (RED)', 'Incomplete with evidence (AMBER)', 'Screening + evidence ready (GREEN)', 'Historical entries (GREY)']])
    for team in teams:
        towers = team['towers']
        summary.append([team['name'], len(towers), sum(t['archive_images'] == 0 for t in towers),
                        sum(t['archive_images'] > 0 and not t['screening_and_evidence_ready'] for t in towers),
                        sum(t['screening_and_evidence_ready'] for t in towers), len(team['outside_assignment'])])
    summary.append([_label('TOTAL'), *[sum(summary.cell(r, c).value for r in range(2, len(teams) + 2)) for c in range(2, 7)]])
    total_row = summary.max_row
    legend_rows = []
    for colour, line, arabic in [
        ('red', 'RED | No evidence: no archived images or field photos by the assigned team. Confirm upload or arrange inspection.', 'الأحمر | لا توجد أدلة: لا توجد صور مؤرشفة أو صور ميدانية من الفريق المكلّف. تأكد من رفع الصور أو رتّب زيارة للفحص.'),
        ('amber', 'AMBER | Incomplete: evidence exists but latest screening or required image categories are incomplete.', 'الأصفر | غير مكتمل: توجد أدلة، لكن الفحص في آخر زيارة أو فئات الصور المطلوبة غير مكتملة.'),
        ('green', 'GREEN | Screening + evidence ready: both coverage checks passed. Final approval and documented exceptions still need review.', 'الأخضر | الفحص والأدلة جاهزة: اكتمل الفحص وتغطية فئات الصور. لا يزال الاعتماد النهائي والاستثناءات الموثّقة بحاجة إلى مراجعة.'),
        ('grey', 'GREY | Historical evidence: outside the current assignment; excluded from assigned totals. Entries may share a tower with another team.', 'الرمادي | أدلة تاريخية: خارج التكليف الحالي ولا تُحتسب ضمن إجمالي الأبراج المكلّف بها الفريق. قد يظهر البرج نفسه في سجلات فريق آخر.'),
        (None, 'An inspection date or visit record alone does not prove a field visit or completion.', 'تاريخ الفحص أو سجل الزيارة وحده لا يثبت تنفيذ زيارة ميدانية أو اكتمال العمل.'),
        (None, 'Tower checklist contains all rows. The four status sheets are separate snapshot lists of the same rows; do not add them again to the full checklist totals.', 'تضم قائمة الأبراج جميع الصفوف. أوراق الحالات الأربع نسخ منفصلة من الصفوف نفسها وقت التصدير؛ لا تضف أعدادها مرة أخرى إلى إجمالي القائمة.'),
        (None, 'Exports include the full selected team checklist; screen search and coverage filters do not restrict it.', 'يشمل التصدير قائمة الفريق المختار كاملة؛ ولا يتأثر بفلاتر البحث أو التغطية على الشاشة.'),
        (None, 'TRUE = Yes; FALSE = No. Readiness is based on the latest inspection; image counts include all saved visits by this team.', 'TRUE = نعم؛ FALSE = لا. تُحدَّد الجاهزية من آخر فحص، بينما يشمل عدد الصور جميع الزيارات المحفوظة لهذا الفريق.'),
        (None, 'Covered categories include saved images and explicit Not required decisions with a recorded reason. No photos are deleted by these decisions.', 'تشمل الفئات المكتملة الصور المحفوظة والقرارات الصريحة بأنها غير مطلوبة مع تسجيل السبب. لا تُحذف الصور بهذه القرارات.'),
    ]:
        summary.append([line + '\n' + arabic])
        summary.merge_cells(start_row=summary.max_row, start_column=1, end_row=summary.max_row, end_column=6)
        summary.row_dimensions[summary.max_row].height = 66
        if colour:
            legend_rows.append((summary.max_row, colour))
    sheet = book.create_sheet(SHEET_TITLES['full'])
    sheet.append([_label(v) for v in ['Team', 'Tower', 'Scope', 'Latest inspection date', 'Visit ID', 'All-time images and field photos',
                  'Installed positions', 'Screened positions', 'Expected category slots', 'Covered category slots',
                  'Missing category slots', 'Screening ready', 'Four-category coverage', 'Screening and evidence ready',
                  'Attention', 'Current assigned team']])
    for team in teams:
        for tower in team['towers']:
            attention = ('No archived evidence' if tower['archive_images'] == 0 else
                         'Screening + evidence ready' if tower['screening_and_evidence_ready'] else 'Incomplete screening or categories')
            sheet.append([team['name'], tower['tower_id'], _label('Current assignment'), tower['inspection_date'],
                          tower['visit_id'], tower['archive_images'], tower['installed_positions'], tower['screened_positions'],
                          tower['expected_categories'], tower['covered_categories'], tower['missing_categories'],
                          tower['screening_ready'], tower['four_category_coverage'], tower['screening_and_evidence_ready'], _label(attention), team['name']])
        for tower in team['outside_assignment']:
            sheet.append([team['name'], tower['tower_id'], _label('Historical evidence outside current assignment'), None, None,
                          tower['archive_images'], *([None] * 8), _label('Historical evidence'), tower['assigned_team_name']])
    if sheet.max_row > 1:
        cells = f'A2:P{sheet.max_row}'
        for formula, fill, colour in [
            ('AND(LEFT($C2,18)="Current assignment",$F2=0)', 'FCE4E6', '9C0006'),
            ('AND(LEFT($C2,18)="Current assignment",$F2>0,$N2=FALSE)', 'FFF2CC', '9C6500'),
            ('AND(LEFT($C2,18)="Current assignment",$F2>0,$N2=TRUE)', 'E2F0D9', '375623'),
            ('LEFT($C2,18)<>"Current assignment"', 'E7E6E6', '595959'),
        ]:
            sheet.conditional_formatting.add(cells, FormulaRule(formula=[formula], fill=PatternFill('solid', fgColor=fill), font=Font(color=colour)))
    headers = [cell.value for cell in sheet[1]]
    status_sheets = {colour: book.create_sheet(SHEET_TITLES[colour]) for colour in STATUS_COLOURS}
    for colour, ws in status_sheets.items():
        ws.append(headers)
        ws.sheet_properties.tabColor = STATUS_COLOURS[colour][1]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        status_sheets[_status(row)].append(row)
    for ws in book:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = 'A2' if ws == summary else 'C2'
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Text identifiers must never become spreadsheet formulas.
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='124B60')
        ws.row_dimensions[1].height = 100
    summary.row_dimensions[total_row].height = 36
    summary.column_dimensions['A'].width = 25
    for col in 'BCDEF':
        summary.column_dimensions[col].width = 25
    for col, colour in [(3, 'red'), (4, 'amber'), (5, 'green'), (6, 'grey')]:
        fill, font = STATUS_COLOURS[colour]
        for row in range(1, total_row + 1):
            summary.cell(row, col).fill = PatternFill('solid', fgColor=fill)
            summary.cell(row, col).font = Font(name='Arial', size=10, color=font, bold=True)
    for row, colour in legend_rows:
        fill, font = STATUS_COLOURS[colour]
        for cell in summary[row]:
            cell.fill = PatternFill('solid', fgColor=fill)
        summary.cell(row, 1).font = Font(name='Arial', size=10, color=font, bold=True)
    for cell in summary[total_row]:
        cell.font = Font(name='Arial', size=10, color=cell.font.color, bold=True)
    widths = [22, 28, 34, 20, 12, 23, 18, 18, 20, 20, 20, 18, 21, 24, 36, 25]
    for ws in [sheet, *status_sheets.values()]:
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 52
            ws.cell(row, 4).number_format = 'yyyy-mm-dd'
            fill, font = STATUS_COLOURS[_status([cell.value for cell in ws[row]])]
            for cell in ws[row]:
                cell.fill = PatternFill('solid', fgColor=fill)
                cell.font = Font(name='Arial', size=10, color=font)
        ws.auto_filter.ref = ws.dimensions
    output = BytesIO()
    book.save(output)
    return output.getvalue()
