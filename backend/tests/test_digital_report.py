import datetime as dt
from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image as PillowImage
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from fastapi import HTTPException

from app.database import Base
from app.models import Image, Position, LineInspectionReport, User
from app.routers import reports
from app.schemas import LineInspectionReportRequest
from app.services.digital_report import archived_evidence_index, export_snapshot, snapshot_rows
from app.services.digital_report import archived_layout, archived_finding_document, NS
from tests.test_oetc_report import _seed


@pytest.fixture
def issued(tmp_path, monkeypatch):
    monkeypatch.setattr(reports.settings, 'images_dir', tmp_path)
    monkeypatch.setattr(reports.settings, 'reports_dir', tmp_path / 'reports')
    engine = create_engine('sqlite:///:memory:')
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        team, tower, _ = _seed(db)
        pos = db.query(Position).first()
        user = User(username='digital-admin', role='admin', hashed_password='x')
        db.add(user)
        for seq, color in enumerate(['red', 'green'], 1):
            name = f'{color}.jpg'
            PillowImage.new('RGB', (40, 30), color).save(tmp_path / name)
            db.add(Image(position_id=pos.id, image_type='TH Full', sequence=seq, file_path=name, include_in_report=True))
        db.commit()
        reports.oetc_line_report(LineInspectionReportRequest(team_id=team.id, start_date=dt.date(2026, 9, 1), end_date=dt.date(2026, 9, 30), report_number='DIGITAL-1'), db, user)
        record = db.query(LineInspectionReport).first()
        path = reports.settings.reports_dir / record.file_path
        yield db, user, record, path, pos


def test_archived_photos_survive_live_replacement(issued):
    db, user, record, path, pos = issued
    index = archived_evidence_index(path, record.inspection_snapshot)
    key = snapshot_rows(record.inspection_snapshot)[0]['key']
    assert set(index) == {f'{key}:0', f'{key}:1'}
    pos.inspector_notes = 'Changed after issue'
    for image in pos.images:
        (reports.settings.images_dir / image.file_path).unlink()
    db.commit()
    response = reports.digital_report_evidence_file(record.id, f'{key}:0', db, user)
    with PillowImage.open(BytesIO(response.body)) as photo:
        assert photo.getpixel((0, 0))[0] > 240
    assert 'Changed after issue' not in str(record.inspection_snapshot)


def test_exports_use_frozen_selected_rows_and_embed_original_photos(issued):
    _, _, record, path, _ = issued
    keys = [snapshot_rows(record.inspection_snapshot)[0]['key']]
    data = export_snapshot(record.inspection_snapshot, record.report_number, record.created_at, keys, 'docx', path)
    doc = Document(BytesIO(data))
    assert len(doc.inline_shapes) == 2
    assert 'not a newly approved report' in '\n'.join(p.text for p in doc.paragraphs)
    data = export_snapshot(record.inspection_snapshot, record.report_number, record.created_at, keys, 'xlsx')
    book = load_workbook(BytesIO(data))
    assert book['Findings'].max_row == 2
    assert book['Findings'].auto_filter.ref
    assert book['Report source']['B1'].value == 'DIGITAL-1'


def test_spreadsheet_notes_are_text_and_selection_is_scoped():
    snapshot = {'team_name': '=1+1', 'visits': [{'id': 1, 'tower': 'T-1', 'positions': [{'id': 2, 'inspector_notes': '=HYPERLINK("https://example.test")'}]}]}
    data = export_snapshot(snapshot, '=1+1', '', ['1:2'], 'xlsx')
    book = load_workbook(BytesIO(data))
    notes = next(cell.column for cell in book['Findings'][1] if cell.value == 'Inspector Notes')
    assert book['Findings'].cell(2, notes).data_type == 's'
    assert book['Report source']['B1'].data_type == 's'
    with pytest.raises(ValueError):
        export_snapshot(snapshot, 'R1', '', ['999:2'], 'xlsx')


def test_digital_permissions_match_original_report(issued):
    db, _, record, _, _ = issued
    outsider = User(username='outsider', role='team_leader', team_id=999, hashed_password='x')
    with pytest.raises(HTTPException) as exc:
        reports.digital_report_evidence(record.id, db, outsider)
    assert exc.value.status_code == 403
    member = User(username='member', role='team_member', hashed_password='x')
    with pytest.raises(HTTPException) as exc:
        reports.digital_report_export(record.id, reports.DigitalReportExport(kind='xlsx', finding_keys=['1:1']), db, member)
    assert exc.value.status_code == 403


def test_corrupt_or_missing_original_has_no_guessed_images(tmp_path):
    path = tmp_path / 'bad.docx'
    path.write_bytes(b'not a zip')
    assert archived_evidence_index(path, {}) == {}
    assert archived_evidence_index(None, {}) == {}


def test_ambiguous_finding_labels_do_not_borrow_photographs(issued):
    import copy
    _, _, record, path, _ = issued
    snapshot = copy.deepcopy(record.inspection_snapshot)
    duplicate = copy.deepcopy(snapshot['visits'][0])
    duplicate['id'] = 999
    snapshot['visits'].append(duplicate)
    assert archived_evidence_index(path, snapshot) == {}


def test_client_can_export_but_older_report_never_uses_live_data(issued):
    db, _, record, _, _ = issued
    client = User(username='customer', role='client', hashed_password='x')
    keys = [snapshot_rows(record.inspection_snapshot)[0]['key']]
    response = reports.digital_report_export(record.id, reports.DigitalReportExport(kind='xlsx', finding_keys=keys), db, client)
    assert response.body.startswith(b'PK')
    record.inspection_snapshot = None
    db.commit()
    with pytest.raises(HTTPException) as exc:
        reports.digital_report_export(record.id, reports.DigitalReportExport(kind='xlsx', finding_keys=keys), db, client)
    assert exc.value.status_code == 404


def test_replica_copies_original_table_and_image_bytes_without_reformatting(issued):
    from zipfile import ZipFile
    from lxml import etree
    _, _, record, path, _ = issued
    layout = archived_layout(path, record.inspection_snapshot)
    assert len(layout) == 2 and all(item['key'] for item in layout)
    key = layout[0]['key']
    data = archived_finding_document(path, record.inspection_snapshot, [key])
    with ZipFile(path) as original, ZipFile(BytesIO(data)) as replica:
        source = etree.fromstring(original.read('word/document.xml')).xpath('./w:body/w:tbl', namespaces=NS)[layout[0]['table_index']]
        copied = etree.fromstring(replica.read('word/document.xml')).xpath('./w:body/w:tbl', namespaces=NS)
        assert len(copied) == 1
        assert etree.tostring(source, method='c14n', exclusive=True) == etree.tostring(copied[0], method='c14n', exclusive=True)
        assert original.read('word/styles.xml') == replica.read('word/styles.xml')
        for name in replica.namelist():
            if name.startswith('word/media/'):
                assert replica.read(name) == original.read(name)
    assert len(Document(BytesIO(data)).inline_shapes) == 2


def test_replica_preserves_selected_order_and_rejects_guessed_tables(issued):
    import copy
    _, _, record, path, _ = issued
    layout = archived_layout(path, record.inspection_snapshot)
    data = archived_finding_document(path, record.inspection_snapshot, [item['key'] for item in reversed(layout)])
    tables = Document(BytesIO(data)).tables
    assert 'No.2' in tables[0].cell(0, 0).text
    assert 'No.1' in tables[1].cell(0, 0).text
    ambiguous = copy.deepcopy(record.inspection_snapshot)
    duplicate = copy.deepcopy(ambiguous['visits'][0]); duplicate['id'] = 999
    ambiguous['visits'].append(duplicate)
    with pytest.raises(ValueError):
        archived_finding_document(path, ambiguous, [layout[0]['key']])


def test_overview_preserves_original_summary_and_approval_tables(issued):
    _, _, record, path, _ = issued
    overview = Document(BytesIO(archived_finding_document(path, record.inspection_snapshot, [], overview=True)))
    text = ' '.join(cell.text for table in overview.tables for row in table.rows for cell in row.cells)
    assert 'Prepared By' in text and 'Overall Condition' in text and 'Remarks' in text
    assert 'Insulator Details' not in text


def test_measurements_preserve_native_rows_and_selected_order(issued):
    from zipfile import ZipFile
    from lxml import etree
    _, _, record, path, _ = issued
    layout = archived_layout(path, record.inspection_snapshot)
    assert all('measurement' in item for item in layout)
    keys = [item['key'] for item in reversed(layout)]
    data = archived_finding_document(path, record.inspection_snapshot, keys, measurements=True)
    with ZipFile(path) as source, ZipFile(BytesIO(data)) as output:
        original = etree.fromstring(source.read('word/document.xml')).xpath('./w:body/w:tbl', namespaces=NS)
        copied = etree.fromstring(output.read('word/document.xml')).xpath('./w:body/w:tbl/w:tr', namespaces=NS)[1:]
        assert len(copied) == len(keys)
        for row, item in zip(copied, reversed(layout)):
            native = original[item['measurement_table_index']].findall('w:tr', NS)[item['measurement_row_index']]
            assert etree.tostring(row, method='c14n', exclusive=True) == etree.tostring(native, method='c14n', exclusive=True)
