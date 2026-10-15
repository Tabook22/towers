from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import Base
from app.models import Position, Visit
from app.schemas import LineInspectionReportRequest
from app.services.oetc_report import render_oetc_line_report_docx
from tests.test_oetc_report import _seed


@pytest.mark.parametrize('mount,direction,gs_side,expected', [
    ('Tension', 'Shahaon', 'Tower side', 'If tension (Direction: Shahaon)'),
    ('Tension', 'Ittin', None, 'If tension (Direction: Ittin)'),
    ('Tension', None, 'Tower side', 'If tension (Direction: Not recorded)'),
    ('Suspension', 'Shahaon', 'Tower side', 'If tension (Direction: —)'),
])
def test_configuration_prints_tension_direction_not_gs_side(mount, direction, gs_side, expected):
    engine = create_engine('sqlite:///:memory:')
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        team, _, _ = _seed(db)
        position = db.query(Position).first()
        position.mount_type = mount
        position.direction = direction
        position.gs_side = gs_side
        position.inspector_notes = 'Direction QA'
        db.commit()
        visit = db.get(Visit, position.visit_id)
        payload = LineInspectionReportRequest(team_id=team.id, start_date=visit.inspection_date, end_date=visit.inspection_date, report_number='DIRECTION-QA')
        data = render_oetc_line_report_docx(team, [visit], payload)
        doc = Document(BytesIO(data))
        lines = [paragraph.text for table in doc.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs]
        assert expected in lines
        assert not any('GS side)' in line for line in lines)
    engine.dispose()
