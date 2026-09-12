"""Excel export of the tower catalog — same columns as import, plus assigned team / active."""
import io

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import Base
from app.models import Team, Tower
from app.services.tower_import import EXPORT_HEADER, HEADER_ROW, export_towers_workbook, import_towers_from_excel


def test_export_includes_all_towers_and_can_reimport():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        team = Team(name="Alpha")
        db.add(team)
        db.flush()
        db.add_all(
            [
                Tower(
                    tower_id="T50",
                    voltage="132 kV",
                    area="Dufar",
                    latitude=17.0,
                    longitude=54.0,
                    assigned_team_id=team.id,
                ),
                Tower(tower_id="T51", area="Dufar", is_active=False),
            ]
        )
        db.commit()
        towers = db.query(Tower).order_by(Tower.tower_id).all()
        for t in towers:
            t.assigned_team  # load

        raw = export_towers_workbook(towers)
        wb = load_workbook(io.BytesIO(raw), data_only=True)
        ws = wb["Towers"]
        headers = [c.value for c in ws[1]]
        assert headers[: len(HEADER_ROW)] == HEADER_ROW
        assert headers == EXPORT_HEADER
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert [r[0] for r in rows] == ["T50", "T51"]
        t50 = rows[0]
        assert t50[1] == "132 kV"
        assert t50[3] == "Dufar"
        assert t50[7] == 17.0
        assert t50[8] == 54.0
        assert t50[10] == "Alpha"
        assert t50[11] == "Yes"
        assert rows[1][11] == "No"

        result = import_towers_from_excel(db, raw)
        assert result["created"] == 0
        assert result["updated"] == 2
        assert result["warnings"] == []
