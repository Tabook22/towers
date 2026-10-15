from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.routers.teams import tower_reconciliation_xlsx
from app.services.tower_reconciliation_export import build_reconciliation_workbook, SHEET_TITLES
from tests.test_tower_reconciliation import data


def test_export_counts_historical_scope_and_safe_text():
    def tower(images, ready):
        return dict(tower_id='=SUM(1,2)', inspection_date=None, visit_id=None, archive_images=images,
                    installed_positions=12, screened_positions=0, expected_categories=48, covered_categories=0,
                    missing_categories=48, screening_ready=ready, four_category_coverage=ready,
                    screening_and_evidence_ready=ready)
    team = dict(name='+Alpha', towers=[tower(0, False), tower(10, False), tower(48, True)],
                outside_assignment=[dict(tower_id='Old', archive_images=66, assigned_team_name='Bravo')])
    book = load_workbook(BytesIO(build_reconciliation_workbook([team])))
    assert list(book[SHEET_TITLES['summary']].values)[1] == ('+Alpha', 3, 1, 1, 1, 1)
    sheet = book[SHEET_TITLES['full']]
    assert sheet.max_row == 5
    assert sheet['A2'].data_type == sheet['B2'].data_type == 's'
    assert sheet['B2'].value == '=SUM(1,2)'
    assert sheet['F2'].value == 0
    assert sheet['N2'].value is False
    assert sheet['P5'].value == 'Bravo'
    assert sheet['G5'].value is None
    rules = sheet.conditional_formatting[next(iter(sheet.conditional_formatting))]
    assert [r.formula[0] for r in rules] == ['AND(LEFT($C2,18)="Current assignment",$F2=0)', 'AND(LEFT($C2,18)="Current assignment",$F2>0,$N2=FALSE)', 'AND(LEFT($C2,18)="Current assignment",$F2>0,$N2=TRUE)', 'LEFT($C2,18)<>"Current assignment"']
    assert [r.dxf.fill.fgColor.rgb for r in rules] == ['00FCE4E6', '00FFF2CC', '00E2F0D9', '00E7E6E6']
    assert sheet.auto_filter.ref == 'A1:P5'
    for name, colour, source in [('red', '00FCE4E6', 2), ('amber', '00FFF2CC', 3),
                                 ('green', '00E2F0D9', 4), ('grey', '00E7E6E6', 5)]:
        ws = book[SHEET_TITLES[name]]
        assert ws.max_row == 2
        assert list(ws.values)[1] == list(sheet.values)[source - 1]
        assert all(cell.fill.fgColor.rgb == colour for cell in ws[2])
        assert all(cell.fill.fgColor.rgb == colour for cell in sheet[source])
        assert ws['B2'].data_type == 's'
    summary = book[SHEET_TITLES['summary']]
    assert summary['C2'].fill.fgColor.rgb == '00FCE4E6'
    assert summary['A4'].fill.fgColor.rgb == '00FCE4E6'
    assert summary['A1'].value == 'Team\nالفريق'
    assert '\nالأحمر' in summary['A4'].value
    for ws in book:
        assert all('\n' in cell.value for cell in ws[1])
    assert sheet['C2'].value == 'Current assignment\nضمن التكليف الحالي'
    assert '\n' in sheet['O2'].value


def test_empty_status_sheets_still_have_filterable_headers():
    book = load_workbook(BytesIO(build_reconciliation_workbook([])))
    assert len(book.sheetnames) == 6
    for name in book.sheetnames[1:]:
        assert book[name].max_row == 1
        assert book[name].auto_filter.ref == 'A1:P1'


def test_export_honours_existing_role_and_team_scope(data):
    db, a, b, towers, admin = data
    from app.models import User
    leader = User(username='lead-export', hashed_password='x', role='team_leader', team_id=a.id)
    response = tower_reconciliation_xlsx(team_id=a.id, db=db, user=leader)
    sheet = load_workbook(BytesIO(response.body))[SHEET_TITLES['full']]
    assert {r[0] for r in list(sheet.values)[1:]} == {a.name}
    with pytest.raises(HTTPException) as exc:
        tower_reconciliation_xlsx(team_id=b.id, db=db, user=leader)
    assert exc.value.status_code == 404
    with pytest.raises(HTTPException) as exc:
        tower_reconciliation_xlsx(team_id=None, db=db, user=User(role='client'))
    assert exc.value.status_code == 403
